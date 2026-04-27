
import openpyxl as xl
import cantools
import sys

SIGNAL_BIBLE_FILE = "Signal_Bible_DP17.xlsx"
SIGNAL_SHEET = "Signal_Bible"

# Colonne Excel
NAME_COL = 1
UNIT_COL = 5
MINMAX_COL = 10
OFFSET_COL = 14
SCALE_COL = 15

START_ROW = 4
END_ROW = 526


class SignalData:
    name: str
    scale: float | None
    offset: float | None
    min_value: float | None
    max_value: float | None
    unit: str | None
    source: str | None


def get_from_CAN0():
    signal_data_arr = []

    for msg in db.messages:
        # message_names.append(msg.name)
        for signal in msg.signals:
            current_signal = SignalData()
            current_signal.name = signal.name
            current_signal.scale = signal.scale
            current_signal.offset = signal.offset
            current_signal.min_value = signal.minimum
            current_signal.max_value = signal.maximum
            current_signal.unit = signal.unit
            current_signal.source = signal.receivers
            signal_data_arr.append(current_signal)
    return signal_data_arr

def open_signal_bible(file_name, sheet_name):
    wb = xl.load_workbook(file_name)
    ws = wb[sheet_name]
    return ws

def access_signal_name(work_sheet):
    signal_names_array = []
    for signal_cell_tu in work_sheet.iter_cols(min_col=NAME_COL, max_col=NAME_COL, min_row=START_ROW, max_row=END_ROW):
        for signal_name in signal_cell_tu:
            signal_names_array.append(signal_name.value)
    return signal_names_array

def signal_not_in_sheet(signals_in_bible, signal_CAN0):
    for signal_in_CAN in signal_CAN0:
        if signal_in_CAN.name not in signals_in_bible:
            print("not found: " + signal_in_CAN)

def signal_not_in_can(signal_in_bible, signal_CAN0):
    for signal_in_CAN in signal_CAN0:
        if signal_in_bible != signal_in_CAN.name:
            continue
        return
    print("not found: " + signal_in_bible)

def is_equal_or_null(signal_from_CAN, signal_sheet_name, signal_sheet_scale, signal_sheet_offset, signal_sheet_min_max, signal_sheet_unit):
    for signal_in_CAN in signal_from_CAN:
        if signal_in_CAN.name in signal_sheet_name:
            index_in_sheet = signal_sheet_name.index(signal_in_CAN.name)
            if (signal_sheet_unit[index_in_sheet] != None and signal_in_CAN.unit != None):
                if (signal_sheet_unit[index_in_sheet] != signal_in_CAN.unit):
                    print(signal_in_CAN.name + " unit error")
                    print(f"in .dbc {signal_in_CAN.unit} VS in signal bible {signal_sheet_unit[index_in_sheet]}")
            if (signal_sheet_scale[index_in_sheet] != None and signal_in_CAN.scale != None):
                if (signal_sheet_scale[index_in_sheet] != signal_in_CAN.scale):
                    print(signal_in_CAN.name + " scale error")
                    print(f"in .dbc {signal_in_CAN.scale} VS in signal bible {signal_sheet_scale[index_in_sheet]}")
            if (signal_sheet_offset[index_in_sheet] != None and signal_in_CAN.offset != None):
                if (signal_sheet_offset[index_in_sheet] != signal_in_CAN.offset):
                    print(signal_in_CAN.name + " offset error")
                    print(f"in .dbc {signal_in_CAN.offset} VS in signal bible {signal_sheet_offset[index_in_sheet]}")
            if (signal_sheet_min_max[index_in_sheet] != None and signal_sheet_min_max[index_in_sheet] != "TBD" and signal_in_CAN.min_value != None and signal_in_CAN.max_value != None):
                CAN_min_max = f"[{signal_in_CAN.min_value} {signal_in_CAN.max_value}]"
                CAN_min_max_second = f"{{{signal_in_CAN.min_value}, {signal_in_CAN.max_value}}}"
                CAN_min_max_third = f"[{signal_in_CAN.min_value} +{signal_in_CAN.max_value}]"
                if (signal_sheet_min_max[index_in_sheet] != CAN_min_max and signal_sheet_min_max[index_in_sheet] != CAN_min_max_second and signal_sheet_min_max[index_in_sheet] != CAN_min_max_third):
                    print(signal_in_CAN.name + " min, max  error")
                    print(f"in .dbc {repr(CAN_min_max)} VS in signal bible {repr(signal_sheet_min_max[index_in_sheet])}")
            #if (signal_in_CAN.source):
                #if (not signal_sheet_name[index_in_sheet].startswith(signal_in_CAN.source[0])):
                    #print(signal_in_CAN.name + " source error")





def access_signal_scale(work_sheet):
    signal_scale_array = []
    for signal_cell_tu in work_sheet.iter_cols(min_col=SCALE_COL, max_col=SCALE_COL, min_row=4, max_row=526):
        for signal_scale in signal_cell_tu:
            signal_scale_array.append(signal_scale.value)
    return signal_scale_array

def access_signal_offset(work_sheet):
    signal_offset_array = []
    for signal_cell_tu in work_sheet.iter_cols(min_col=OFFSET_COL, max_col=OFFSET_COL, min_row=4, max_row=526):
        for signal_offset in signal_cell_tu:
            signal_offset_array.append(signal_offset.value)
    return signal_offset_array

def access_signal_min_max(work_sheet):
    signal_min_max_array = []
    for signal_cell_tu in work_sheet.iter_cols(min_col=MINMAX_COL, max_col=MINMAX_COL, min_row=4, max_row=526):
        for signal_min_max in signal_cell_tu:
            signal_min_max_array.append(signal_min_max.value)
    return signal_min_max_array

def access_signal_unit(work_sheet):
    signal_unit_array = []
    for signal_cell_tu in work_sheet.iter_cols(min_col=UNIT_COL, max_col=UNIT_COL, min_row=4, max_row=526):
        for signal_unit in signal_cell_tu:
            signal_unit_array.append(signal_unit.value)
    return signal_unit_array

if __name__ == '__main__':
    if len(sys.argv) > 1:
        dbc_file = sys.argv[1]
        try:
            db = cantools.database.load_file(dbc_file)
        except FileNotFoundError:
            print(f"Errore: Il file {dbc_file} non esiste.")
        signal_sheet = open_signal_bible(SIGNAL_BIBLE_FILE,SIGNAL_SHEET )
        signal_bible_names = access_signal_name(signal_sheet)
        signal_bible_offset = access_signal_offset(signal_sheet)
        signal_bible_scale = access_signal_scale(signal_sheet)
        signal_bible_min_max = access_signal_min_max(signal_sheet)
        signal_bible_unit = access_signal_unit(signal_sheet)
        signal_CAN0 = get_from_CAN0()
        is_equal_or_null(signal_CAN0,signal_bible_names,signal_bible_scale,signal_bible_offset,signal_bible_min_max, signal_bible_unit)
    else:
        print("errore: nessun file inserito")









